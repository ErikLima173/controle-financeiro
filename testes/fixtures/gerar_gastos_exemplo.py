"""
Gera testes/fixtures/gastos-exemplo.xlsx: mesmo layout da planilha Gastos.xlsx do Erik,
com valores FICTÍCIOS (o repositório é público; a planilha real não pode ser versionada).

    Gastos Nathy    tabela Gastos_Mensais A1:Q13   O = SOMA(B:N), Q = P - O
    Gastos vini     tabela Gasto_Mensal   A1:K14   J = SOMA(B:H), K = I - J
    Total finanças  tabela Monthly_Finances_Review A1:D13, soma as duas abas acima
    Gastos (oculta) histórico 2025/2026 em blocos, com células mescladas e meses com espaço no fim

As fórmulas saem com o valor já calculado (<v>), como num arquivo salvo pelo Excel ou pelo
Google Planilhas, para que quem lê só os valores veja os totais.

Junto sai saida.esperado.json com os totais de cada mês (por pessoa e somados), para os testes
conferirem o que o sistema mostra depois de importar.

Uso (precisa do openpyxl: pip install openpyxl):
    python3 testes/fixtures/gerar_gastos_exemplo.py            # regrava gastos-exemplo.xlsx ao lado deste arquivo
    python3 testes/fixtures/gerar_gastos_exemplo.py outra.xlsx
O .gitignore bloqueia *.xlsx (para planilhas pessoais não subirem por engano); esta foi adicionada com
git add -f e, por já estar no repositório, as próximas versões entram normalmente.
"""
import json
import random
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

MOEDA = '[$R$ -416]#,##0.00'
MESES = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
         'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
MESES_PREENCHIDOS = 9  # janeiro a setembro com valores; outubro a dezembro em branco

rnd = random.Random(173)
calculado = {}  # (aba, célula) -> valor da fórmula


def centavos(a, b):
    return round(rnd.uniform(a, b), 2)


def tabela(ws, nome, ref):
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showFirstColumn=True)
    ws.add_table(t)


def cabecalho(ws, nomes):
    for c, nome in enumerate(nomes, 1):
        ws.cell(1, c, nome).font = Font(bold=True)


def aba_nathy(wb):
    ws = wb.active
    ws.title = 'Gastos Nathy'
    cabecalho(ws, ['Mes', 'Cartão', 'Academia', 'Óculos', 'máquina', 'PC', 'Areia Gatos', 'Ração Athena',
                   'Ração Gatos', 'Comida', 'Compras', 'Jogos', 'Gasolina', 'Outros', 'Total Gasto',
                   'Salário', 'Total restante'])
    faixas = {  # coluna: (mín, máx, chance de ficar em branco)
        2: (350, 1400, 0), 3: (99.9, 99.9, 0), 4: (0, 0, 1), 5: (60, 60, 0.5), 6: (0, 0, 1),
        7: (35, 70, 0.2), 8: (90, 140, 0), 9: (80, 120, 0), 10: (300, 750, 0), 11: (0, 420, 0.4),
        12: (0, 180, 0.5), 13: (150, 320, 0), 14: (0, 200, 0.3),
    }
    for i, mes in enumerate(MESES):
        r = i + 2
        ws.cell(r, 1, mes)
        soma = 0.0
        if i < MESES_PREENCHIDOS:
            if i == 4:
                faixas[4] = (480, 480, 0)  # óculos só em maio
            elif i == 6:
                faixas[6] = (3299.9, 3299.9, 0)  # PC em julho
            for c, (a, b, branco) in faixas.items():
                if rnd.random() < branco:
                    continue
                v = centavos(a, b)
                ws.cell(r, c, v).number_format = MOEDA
                soma += v
            faixas[4] = (0, 0, 1)
            faixas[6] = (0, 0, 1)
            salario = 3450.0
            ws.cell(r, 16, salario).number_format = MOEDA
        else:
            salario = 0.0
        ws.cell(r, 15, f'=SUM(B{r}:N{r})').number_format = MOEDA
        ws.cell(r, 17, f'=P{r}-O{r}').number_format = MOEDA
        calculado[(ws.title, f'O{r}')] = round(soma, 2)
        calculado[(ws.title, f'Q{r}')] = round(salario - soma, 2)
    tabela(ws, 'Gastos_Mensais', 'A1:Q13')
    ws.freeze_panes = 'B2'
    return ws


def aba_vini(wb):
    ws = wb.create_sheet('Gastos vini')
    cabecalho(ws, ['Mês', 'Faculdade', 'Internet', 'Seguro Carro', 'academia', 'Aluguel', 'Água',
                   'Gastos extras', 'Salário', 'Total Gasto', 'Total Restante'])
    for i, mes in enumerate(MESES):
        r = i + 2
        ws.cell(r, 1, mes)
        soma = 0.0
        if i < MESES_PREENCHIDOS:
            valores = {2: 450.0, 3: 99.9, 4: 187.35, 5: 89.9 if i % 3 else None, 6: 1150.0,
                       7: centavos(38, 72), 8: centavos(0, 650) if i % 2 else None}
            for c, v in valores.items():
                if v is None:
                    continue
                ws.cell(r, c, v).number_format = MOEDA
                soma += v
            salario = 2980.0 if i != 5 else 4470.0  # junho com 13º adiantado
            ws.cell(r, 9, salario).number_format = MOEDA
        else:
            salario = 0.0
        ws.cell(r, 10, f'=SUM(B{r}:H{r})').number_format = MOEDA
        ws.cell(r, 11, f'=I{r}-J{r}').number_format = MOEDA
        calculado[(ws.title, f'J{r}')] = round(soma, 2)
        calculado[(ws.title, f'K{r}')] = round(salario - soma, 2)
    tabela(ws, 'Gasto_Mensal', 'A1:K14')  # como no original: uma linha vazia sobrando no fim
    ws.freeze_panes = 'B2'
    return ws


def aba_total(wb):
    ws = wb.create_sheet('Total finanças')
    cabecalho(ws, ['mes', 'Total Recebido', 'Total Gasto', 'Total Restante'])
    for i, mes in enumerate(MESES):
        r = i + 2
        ws.cell(r, 1, mes.lower())
        ws.cell(r, 2, f"=SUM('Gastos Nathy'!P{r}+'Gastos vini'!I{r})").number_format = MOEDA
        ws.cell(r, 3, f"=SUM('Gastos Nathy'!O{r}+'Gastos vini'!J{r})").number_format = MOEDA
        ws.cell(r, 4, f'=B{r}-C{r}').number_format = MOEDA
        rec = (wb['Gastos Nathy'].cell(r, 16).value or 0) + (wb['Gastos vini'].cell(r, 9).value or 0)
        gas = calculado[('Gastos Nathy', f'O{r}')] + calculado[('Gastos vini', f'J{r}')]
        calculado[(ws.title, f'B{r}')] = round(rec, 2)
        calculado[(ws.title, f'C{r}')] = round(gas, 2)
        calculado[(ws.title, f'D{r}')] = round(rec - gas, 2)
    tabela(ws, 'Monthly_Finances_Review', 'A1:D13')
    ws.freeze_panes = 'A2'
    return ws


def aba_historico(wb):
    ws = wb.create_sheet('Gastos')
    ws.sheet_state = 'hidden'
    ws['A1'] = 'Mes'
    ws['B1'] = 2025
    ws.merge_cells('A1:A2')
    ws.merge_cells('B1:Q1')
    nomes = ['aluguel', 'luz ', 'agua', 'internet', 'faculdade', 'youtube', 'tvbox', 'parcela tv ',
             'parcela ps5', 'Gastos extras', 'ração porquinho', 'ração athena', 'ração gato',
             'total gasto', 'Valor recebido', 'valor restante ']
    for c, nome in enumerate(nomes, 2):
        ws.cell(2, c, nome)
    fixos = {2: 1150.0, 5: 99.9, 6: 450.0, 7: 24.9, 8: 35.0, 9: 180.0, 10: 275.0}
    meses_2025 = ['março', 'abril', 'maio', 'junho ', 'julho ', 'agosto', 'setembro', 'outubro',
                  'novembro', 'dezembro']
    for i, mes in enumerate(meses_2025):
        r = i + 3
        ws.cell(r, 1, mes).number_format = MOEDA
        valores = dict(fixos)
        if i < 6:
            valores[3] = centavos(70, 130)  # luz
        if i < 5:
            valores[4] = centavos(35, 62)  # água
        if i in (1, 2, 3):
            valores[11] = centavos(90, 1200)  # gastos extras
        if i == 9:
            del valores[9]  # última parcela da TV foi em novembro
        for c, v in valores.items():
            ws.cell(r, c, v).number_format = MOEDA
        ws.cell(r, 15, f'=SUM(B{r}:N{r})').number_format = MOEDA
        ws.cell(r, 16, 3200.0).number_format = MOEDA
        ws.cell(r, 17, f'=P{r}-O{r}').number_format = MOEDA
        soma = round(sum(valores.values()), 2)
        calculado[(ws.title, f'O{r}')] = soma
        calculado[(ws.title, f'Q{r}')] = round(3200.0 - soma, 2)
    ws['A13'] = 2026
    ws.merge_cells('A13:Q13')
    for i, mes in enumerate(['Janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho ', 'julho ',
                             'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']):
        r = i + 14
        ws.cell(r, 1, mes)
        for c, v in {2: 1150.0, 5: 99.9, 7: 24.9, 8: 35.0}.items():
            ws.cell(r, c, v).number_format = MOEDA
        if i == 0:
            ws.cell(r, 10, 275.0).number_format = MOEDA  # última parcela do PS5
    return ws


def nomes_definidos(wb):
    for nome, ref in [('Totalgastonathy', "'Gastos Nathy'!$O$1:$O$13"),
                      ('Totalrestantenathy', "'Gastos Nathy'!$Q$1:$Q$13"),
                      ('Totalgastovini', "'Gastos vini'!$J$1:$J$13"),
                      ('Totalrestantevini', "'Gastos vini'!$K$1:$K$13")]:
        wb.defined_names[nome] = DefinedName(nome, attr_text=ref)


def larguras(wb):
    for ws in wb.worksheets:
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(c)].width = 15


def gravar_valores_calculados(caminho, abas):
    """openpyxl grava fórmulas sem resultado; preenche o <v> de cada fórmula com o valor calculado."""
    tmp = caminho.with_suffix('.tmp')
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            m = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', item.filename)
            if m:
                aba = abas[int(m.group(1)) - 1]
                xml = dados.decode('utf8')

                def preencher(mt):
                    ref = mt.group(1)
                    v = calculado.get((aba, ref))
                    if v is None:
                        raise SystemExit(f'fórmula sem valor calculado: {aba}!{ref}')
                    return f'{mt.group(0)[:-len(mt.group(3))]}<v>{v:g}</v>' if mt.group(3) else mt.group(0)

                xml, n = re.subn(r'<c r="([A-Z]+\d+)"[^>]*>(<f>[^<]*</f>)(<v\s*/>|<v></v>)?', preencher, xml)
                dados = xml.encode('utf8')
            zout.writestr(item, dados)
    tmp.replace(caminho)


def esperado(wb):
    nathy, vini = wb['Gastos Nathy'], wb['Gastos vini']
    meses = []
    for i, mes in enumerate(MESES):
        r = i + 2
        meses.append({
            'mes': mes,
            'nathy': {'gasto': calculado[('Gastos Nathy', f'O{r}')], 'salario': nathy.cell(r, 16).value or 0,
                      'restante': calculado[('Gastos Nathy', f'Q{r}')]},
            'vini': {'gasto': calculado[('Gastos vini', f'J{r}')], 'salario': vini.cell(r, 9).value or 0,
                     'restante': calculado[('Gastos vini', f'K{r}')]},
            'total': {'recebido': calculado[('Total finanças', f'B{r}')],
                      'gasto': calculado[('Total finanças', f'C{r}')],
                      'restante': calculado[('Total finanças', f'D{r}')]},
        })
    soma = lambda chave: round(sum(m['total'][chave] for m in meses), 2)
    return {
        'aviso': 'Valores fictícios gerados por gerar_gastos_exemplo.py; não são dados reais.',
        'meses': meses,
        'ano': {'recebido': soma('recebido'), 'gasto': soma('gasto'), 'restante': soma('restante')},
        'historico_2025': [
            {'mes': wb['Gastos'].cell(r, 1).value.strip(), 'gasto': calculado[('Gastos', f'O{r}')],
             'recebido': wb['Gastos'].cell(r, 16).value, 'restante': calculado[('Gastos', f'Q{r}')]}
            for r in range(3, 13)
        ],
        # Bloco de 2026 da aba oculta: só gastos, sem fórmula de total e sem valor recebido.
        'historico_2026': [
            {'mes': wb['Gastos'].cell(r, 1).value.strip(),
             'gasto': round(sum(wb['Gastos'].cell(r, c).value or 0 for c in range(2, 15)), 2)}
            for r in range(14, 26)
        ],
    }


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name('gastos-exemplo.xlsx')
    wb = Workbook()
    aba_nathy(wb)
    aba_vini(wb)
    aba_total(wb)
    aba_historico(wb)
    nomes_definidos(wb)
    larguras(wb)
    wb.save(saida)
    gravar_valores_calculados(saida, [ws.title for ws in wb.worksheets])
    saida.with_suffix('.esperado.json').write_text(
        json.dumps(esperado(wb), ensure_ascii=False, indent=2) + '\n', encoding='utf8')
    print(f'{saida}: {sum(1 for _ in calculado)} fórmulas com valor calculado')


if __name__ == '__main__':
    main()
